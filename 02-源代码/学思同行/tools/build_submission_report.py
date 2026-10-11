"""Build the credential-free 2026-10-11 submission report.

Uses the bundled document runtime. The Markdown is the content authority.
No running .env, database, private textbook, cookie, or model response is read.
"""
from __future__ import annotations

from pathlib import Path
import argparse
from datetime import datetime, timezone
import html
import json
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image as PILImage, ImageDraw, ImageFont
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image,
    PageBreak, KeepTogether,
)
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
FINAL_STEM = "学思同行-项目技术与验收报告-20261011"
BLACK = "000000"
BORDER = "D9D9D9"
FONT = Path("C:/Windows/Fonts/simhei.ttf")
YAHEI = Path("C:/Windows/Fonts/msyh.ttc")


def plain(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", value)
    return value.replace("**", "").replace(chr(96), "")


def blocks(text: str):
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<!-- pagebreak -->":
            yield "pagebreak", None
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [plain(v.strip()) for v in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            yield "table", rows
            continue
        elif match := re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line):
            yield "image", (match.group(1), match.group(2))
        elif line.startswith("#"):
            depth = len(line) - len(line.lstrip("#"))
            yield f"heading{min(depth, 3)}", plain(line.lstrip("#").strip())
        elif line.startswith(("~~~", chr(96) * 3)):
            fence = line[:3]
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith(fence):
                code.append(lines[index])
                index += 1
            yield "code", code
        elif line.startswith(("- ", "* ")):
            yield "bullet", plain(line[2:])
        else:
            yield "paragraph", plain(line)
        index += 1


def architecture_image(path: Path) -> None:
    """A source-based component diagram; it contains no private UI or data."""
    path.parent.mkdir(parents=True, exist_ok=True)
    canvas = PILImage.new("RGB", (2100, 1190), "white")
    draw = ImageDraw.Draw(canvas)
    title_font = ImageFont.truetype(str(YAHEI), 42)
    text_font = ImageFont.truetype(str(YAHEI), 34)
    small_font = ImageFont.truetype(str(YAHEI), 29)
    draw.text((70, 40), "学生问题、实际教材与学习状态连接到同一流程", fill="#000000", font=title_font)

    def box(bounds, title, subtitle):
        draw.rounded_rectangle(bounds, radius=14, outline="#697D75", width=3, fill="#F5F7F5")
        x1, y1, x2, y2 = bounds
        draw.text(((x1+x2)/2, y1+25), title, anchor="mt", font=text_font, fill="#000000")
        for row, line in enumerate(subtitle.splitlines()):
            draw.text(((x1+x2)/2, y1+78+row*43), line, anchor="mt",
                      font=small_font, fill="#26342E")

    def arrow(points, label=None):
        draw.line(points, fill="#466357", width=4, joint="curve")
        x, y = points[-1]
        px, py = points[-2]
        if x != px:
            sign = 1 if x > px else -1
            head = [(x, y), (x-sign*18, y-10), (x-sign*18, y+10)]
        else:
            sign = 1 if y > py else -1
            head = [(x, y), (x-10, y-sign*18), (x+10, y-sign*18)]
        draw.polygon(head, fill="#466357")
        if label:
            mid = points[len(points)//2]
            draw.text((mid[0]+12, mid[1]+10), label, font=small_font, fill="#26342E")

    box((70, 135, 540, 305), "学生学习工作台", "课程 · 对话 · 原文\n练习 · 错题 · 计划")
    box((670, 135, 1220, 305), "FastAPI 业务与鉴权", "Cookie / REST / SSE\n课程和资源所有者检查")
    box((1370, 135, 2020, 305), "SQLite WAL", "资料 · 会话 · 练习 · 计划\n事务 · 草稿版本 · 持久记录")
    box((70, 465, 540, 675), "资料摄取与来源", "PDF / DOCX / TXT / MD\n章节 · 物理页 · 质量 · 进度")
    box((670, 465, 1220, 675), "LangGraph 教学流程", "证据判定 · 任务化教学\n必要时规划与白名单工具")
    box((1370, 465, 2020, 675), "Mimo 模型服务", "教学 · 出题 · 独立审核\n按需单页视觉")
    box((70, 855, 540, 1055), "教材与计算工具", "BM25 / 可选向量 / 相邻页\n二分 · 分页 · 调度等")
    box((670, 855, 1220, 1055), "回答检查与诚实恢复", "协议 · 引用 · 数值 · 事实\n内部失败转通用学习回复")
    box((1370, 855, 2020, 1055), "练习与复习闭环", "来源 · 盲解 · 解析核对\n评分 · 错题 · 计划 · 导出")
    arrow([(540, 220), (670, 220)])
    arrow([(1220, 220), (1370, 220)])
    arrow([(945, 305), (945, 465)])
    arrow([(305, 305), (305, 465)])
    arrow([(540, 565), (670, 565)])
    arrow([(1220, 565), (1370, 565)])
    arrow([(305, 675), (305, 855)])
    arrow([(540, 955), (610, 955), (610, 620), (670, 620)])
    arrow([(945, 675), (945, 855)])
    arrow([(1695, 675), (1695, 745), (1120, 745), (1120, 855)])
    arrow([(2020, 955), (2060, 955), (2060, 220), (2020, 220)])
    draw.text((72, 1120), "实际凭证仅后端读取；文字原文、视觉观察和确定性工具结果分别标注。", font=small_font, fill="#26342E")
    canvas.save(path)


def column_fractions(rows):
    n = len(rows[0])
    if n == 2:
        if rows[0][0] == "配置项":
            return [.46, .54]
        return [.32, .68]
    if n == 3:
        return [.24, .42, .34]
    if n == 4:
        return [.16, .28, .28, .28]
    return [1/n] * n


def word_font(run, size=None, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.color.rgb = RGBColor.from_string(BLACK)
    run.bold = bold
    if size:
        run.font.size = Pt(size)


def word(elements, path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [("Normal", 10.5), ("Title", 23), ("Heading 1", 15.5), ("Heading 2", 12.5), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        # The bundled default Title style includes an accent-colored rule.
        # Formal reports use plain black titles without decorative borders.
        p_props = style._element.get_or_add_pPr()
        for border in list(p_props.findall(qn("w:pBdr"))):
            p_props.remove(border)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.space_before = Pt(12 if name.startswith("Heading") else 0)
    doc.styles["Title"].paragraph_format.space_after = Pt(13)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.3
    doc.styles["Normal"].paragraph_format.widow_control = True
    doc.core_properties.title = "学思同行项目技术与验收报告"
    doc.core_properties.subject = "中兴赛道专业课助教智能体技术和提交材料"
    doc.core_properties.author = "学思同行项目组"
    doc.core_properties.keywords = "专业课助教 教材 学习记录 20261011"
    doc.core_properties.language = "zh-CN"
    doc.core_properties.created = datetime.now(timezone.utc).replace(tzinfo=None)
    doc.core_properties.modified = doc.core_properties.created
    for kind, value in elements:
        if kind.startswith("heading"):
            level = int(kind[-1])
            if level == 1:
                doc.add_paragraph(value, "Title")
            else:
                doc.add_heading(value, level=level-1)
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "table":
            if not value:
                continue
            table = doc.add_table(rows=0, cols=len(value[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            fractions = column_fractions(value)
            for col, fraction in zip(table.columns, fractions):
                col.width = Cm(17*fraction)
            for row_index, contents in enumerate(value):
                row = table.add_row()
                props = row._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    props.append(OxmlElement("w:tblHeader"))
                for col_index, (cell, content) in enumerate(zip(row.cells, contents)):
                    cell.width = Cm(17*fractions[col_index])
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    cell_props = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for edge in ["top", "left", "bottom", "right"]:
                        item = OxmlElement("w:"+edge)
                        item.set(qn("w:val"), "single")
                        item.set(qn("w:sz"), "4")
                        item.set(qn("w:color"), BORDER)
                        borders.append(item)
                    cell_props.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for edge, amount in [("top", 100), ("bottom", 100), ("left", 120), ("right", 120)]:
                        item = OxmlElement("w:"+edge)
                        item.set(qn("w:w"), str(amount))
                        item.set(qn("w:type"), "dxa")
                        margins.append(item)
                    cell_props.append(margins)
                    if row_index == 0 or row_index % 2 == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E8EEF2" if row_index == 0 else "F6F8FA")
                        cell_props.append(shade)
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.2
                    if len(content) < 10:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    word_font(paragraph.add_run(content), 9.5, row_index == 0)
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
        elif kind == "image":
            caption, relative = value
            image_path = OUT / relative
            if not image_path.is_file():
                raise FileNotFoundError("Report image not found: "+str(image_path))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(image_path), width=Cm(17))
            p = doc.add_paragraph(caption)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            for run in p.runs:
                word_font(run, 9)
        elif kind == "code":
            for line in value:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                run = p.add_run(line)
                word_font(run, 9)
                run.font.name = "Consolas"
        else:
            p = doc.add_paragraph(value, "List Bullet" if kind == "bullet" else "Normal")
    header = section.header.paragraphs[0]
    header.text = "学思同行  专业课助教智能体"
    for run in header.runs:
        word_font(run, 8.5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    word_font(footer.add_run("第 "), 8.5)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    word_font(footer.add_run(" 页"), 8.5)
    doc.save(path)
    reopened = Document(path)
    if not reopened.paragraphs or reopened.paragraphs[0].text != "学思同行项目技术与验收报告":
        raise RuntimeError("DOCX reopen title verification failed")
    with ZipFile(path) as zipped:
        if zipped.testzip():
            raise RuntimeError("DOCX ZIP integrity failed")


def pdf(elements, path):
    pdfmetrics.registerFont(TTFont("CN", str(FONT)))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("CNBody", fontName="CN", fontSize=10.2, leading=16,
                              wordWrap="CJK", spaceAfter=7, textColor=colors.black))
    styles.add(ParagraphStyle("CNCell", parent=styles["CNBody"], fontSize=9, leading=14, spaceAfter=0))
    styles.add(ParagraphStyle("CNCode", parent=styles["CNBody"], fontSize=8.6, leading=13, spaceAfter=2))
    styles.add(ParagraphStyle("CNCaption", parent=styles["CNBody"], fontSize=8.8, leading=14,
                              alignment=1, spaceAfter=12))
    for level, size in [(1, 23), (2, 15.5), (3, 12.5)]:
        styles.add(ParagraphStyle(f"CNHeading{level}", fontName="CN", fontSize=size,
                                  leading=size*1.4, textColor=colors.black, spaceBefore=14,
                                  spaceAfter=10, keepWithNext=True, wordWrap="CJK"))
    def paragraph(text, style="CNBody"):
        return Paragraph(html.escape(text).replace("\n", "<br/>"), styles[style])
    flow = []
    width = A4[0] - 2*56.692913
    for kind, value in elements:
        if kind.startswith("heading"):
            flow.append(paragraph(value, "CNHeading"+kind[-1]))
        elif kind == "pagebreak":
            flow.append(PageBreak())
        elif kind == "table":
            if not value:
                continue
            tab = Table([[paragraph(cell, "CNCell") for cell in row] for row in value],
                        colWidths=[width*f for f in column_fractions(value)], repeatRows=1,
                        hAlign="LEFT", splitByRow=1)
            tab.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#E8EEF2")),
                ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F6F8FA")]),
                ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#D9D9D9")),
                ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
                ("TOPPADDING", (0,0), (-1,-1), 7),
                ("BOTTOMPADDING", (0,0), (-1,-1), 7),
                ("LEFTPADDING", (0,0), (-1,-1), 7),
                ("RIGHTPADDING", (0,0), (-1,-1), 7),
            ]))
            flow.extend([tab, Spacer(1, 10)])
        elif kind == "image":
            caption, relative = value
            image_path = OUT / relative
            with PILImage.open(image_path) as img:
                height = width*img.height/img.width
            flow.append(KeepTogether([Image(str(image_path), width=width, height=height),
                                      Spacer(1, 5), paragraph(caption, "CNCaption")]))
        elif kind == "code":
            flow.append(KeepTogether([paragraph(line, "CNCode") for line in value]))
        else:
            flow.append(paragraph(("- " if kind == "bullet" else "")+value))
    def page(canvas, doc):
        canvas.setFont("CN", 8.5)
        canvas.setFillColor(colors.black)
        canvas.drawString(56.7, A4[1]-32, "学思同行  专业课助教智能体")
        canvas.drawRightString(A4[0]-56.7, 29, f"第 {doc.page} 页")
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=56.7, rightMargin=56.7,
                      topMargin=56.7, bottomMargin=48,
                      title="学思同行项目技术与验收报告",
                      author="学思同行项目组").build(flow, onFirstPage=page, onLaterPages=page)
    reader = PdfReader(path)
    if not reader.pages or any(not page.extract_text().strip() for page in reader.pages):
        raise RuntimeError("PDF reopen text verification failed")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=OUT/"submission-report.md")
    parser.add_argument("--output-stem", default=FINAL_STEM)
    args = parser.parse_args()
    architecture_image(OUT/"release-assets"/"architecture-20261011.png")
    elements = list(blocks(args.source.read_text(encoding="utf-8")))
    docx_path = OUT/(args.output_stem+".docx")
    pdf_path = OUT/(args.output_stem+".pdf")
    word(elements, docx_path)
    pdf(elements, pdf_path)
    print(json.dumps({"docx": str(docx_path), "pdf": str(pdf_path),
                      "pdf_pages": len(PdfReader(pdf_path).pages),
                      "source": str(args.source)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
