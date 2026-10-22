from __future__ import annotations

import asyncio
import hashlib
import json
import re
import zipfile
import unicodedata
import bisect
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from .db import dumps, loads, now, uid
from .knowledge_map import derive_textbook_tree, extracted_graph
from .figures import (extract_figure_captions, queue_figure_index, queue_missing_figure_indexes,
                      scan_pdf_figures, store_figure_index)

INGESTION_VERSION = "2026-10-03-normalized-pdf-chapters-v3"

CJK = r"[\u3400-\u9fff]"
PDF_HEADING = re.compile(r"^(?:第[一二三四五六七八九十百0-9]+[章节].{0,65}|\d{1,2}(?:\.\d{1,2}){1,3}\s*.{2,65})$")


def is_pdf_heading(text):
    return bool(PDF_HEADING.match(text)) and bool(re.search(CJK, text)) and not re.search(r"[。，,；;！？!?]|介绍|讲述|阐述|较详细|是一个|主要介绍|本章|\.{3,}|…{2,}|(?:\. ){3,}", text)


class ExtractedPages(list):
    def __init__(self, items=(), sections=None, quality=None, figures=None):
        super().__init__(items)
        self.sections = sections or {}
        self.quality = quality or {}
        self.figures = figures


def normalize_text(text, pdf=False):
    """NFKC compatibility CJK/width cleanup, preserving mathematical notation."""
    protected = lambda char: (0x2070 <= ord(char) <= 0x209F or 0x1D400 <= ord(char) <= 0x1D7FF
                              or char in "²³¹" or unicodedata.category(char) in {"Sm", "Nl"})
    text = "".join(char if protected(char) else unicodedata.normalize("NFKC", char) for char in text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00ad", "")
    lines = []
    for raw in text.splitlines():
        line = re.sub(r"[\t \u00a0]+", " ", raw).strip()
        compact = re.sub(rf"(?<={CJK}) +(?={CJK})", "", line)
        compact = re.sub(r"^第\s+", "第", compact)
        compact = re.sub(r"(?<=章)\s+", "", compact)
        is_heading = is_pdf_heading(compact)
        if pdf and is_heading:
            line = compact
        # Only join an obvious CJK sentence continuation, preserving headings and formula lines.
        previous_heading = bool(lines and is_pdf_heading(lines[-1]))
        if pdf and lines and not is_heading and not previous_heading and re.search(rf"{CJK}$", lines[-1]) and re.match(rf"^{CJK}", line):
            lines[-1] += line
        else:
            lines.append(line)
    return "\n".join(lines).strip()


def usable_outline(reader):
    entries = []

    def visit(items, parents=()):
        previous = parents
        for item in items:
            if isinstance(item, list):
                visit(item, previous)
                continue
            title = normalize_text(str(item.get("/Title", ""))).strip()
            if not title or re.fullmatch(r"(?:text|page|section|bookmark)[ _-]*\d+", title, re.I):
                previous = parents
                continue
            try:
                page = reader.get_destination_page_number(item) + 1
                if page > 0:
                    label = " / ".join((*parents, title))
                    entries.append((page, label))
                    previous = (*parents, title)
            except Exception:
                continue
    try:
        visit(reader.outline)
    except Exception:
        return []
    return sorted(entries, key=lambda item: item[0])


def pdf_sections(pages, outline):
    result, chapter, section = {}, "", "正文"
    outline_pages = [page for page, _ in outline]
    for page, text in pages:
        if outline:
            index = bisect.bisect_right(outline_pages, page) - 1
            if index >= 0:
                section = outline[index][1]
        else:
            initial_section = section
            for index, line in enumerate(text.splitlines()):
                candidate = line.strip()
                if not is_pdf_heading(candidate):
                    continue
                if re.match(r"^第[一二三四五六七八九十百0-9]+章", candidate):
                    chapter, section = candidate, candidate
                else:
                    section = f"{chapter} / {candidate}" if chapter else candidate
                if index == 0:
                    initial_section = section
            result[page] = initial_section
            continue
        result[page] = section
    return result


def extract_pages(path: Path, progress=None):
    """PDF pages are physical; text/DOCX pages are explicitly logical segments."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF 已加密，请上传可读取的版本")
        if len(reader.pages) > 2000:
            raise ValueError("单个文档最多支持 2000 页")
        count = len(reader.pages)
        if progress:
            progress({"stage": "extracting", "current_page": 0, "total_pages": count, "chunk_count": 0, "embedding_percent": 0})
        pages, normalized_chars, figures = [], 0, []
        for i, page in enumerate(reader.pages):
            raw = page.extract_text() or ""
            figures.extend(extract_figure_captions(raw, i + 1))
            normalized_chars += sum(1 for char in raw if unicodedata.normalize("NFKC", char) != char)
            pages.append((i + 1, normalize_text(raw, pdf=True)))
            if progress and (i % 10 == 0 or i + 1 == count):
                progress({"stage": "extracting", "current_page": i + 1, "total_pages": count, "chunk_count": 0, "embedding_percent": 0})
        outline = usable_outline(reader)
        sections = pdf_sections(pages, outline)
        for figure in figures:
            figure["section_path"] = sections.get(figure["physical_page"], "")
        pages = ExtractedPages(pages, sections=sections, quality={"normalization": "NFKC compatibility cleanup; math notation preserved", "normalized_characters": normalized_chars,
                                                               "chapter_source": "PDF outline" if outline else "正文标题提取（待核验）", "outline_entries": len(outline),
                                                               "pdf_outline": [{"page": page, "path": path} for page, path in outline]}, figures=figures)
    elif suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > 120 * 1024 * 1024:
                raise ValueError("DOCX 解压体积过大")
        doc = DocxDocument(path)
        # No renderer can reliably infer Word pagination; stable text segments are labeled.
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        pages = logical_pages(normalize_text(text))
    elif suffix in {".md", ".txt"}:
        raw = path.read_bytes()
        text = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("文本编码无法识别，请转为 UTF-8")
        if "\x00" in text:
            raise ValueError("上传文件不是有效文本")
        pages = logical_pages(normalize_text(text))
    else:
        raise ValueError("文件类型不支持")
    if not any(text.strip() for _, text in pages):
        raise ValueError("未提取到可检索文字；扫描 PDF 请先 OCR 后再上传")
    if not isinstance(pages, ExtractedPages):
        pages = ExtractedPages(pages, quality={"chapter_source": "正文标题提取（待核验）"})
    blank = sum(not text.strip() for _, text in pages)
    short = sum(0 < len(text.strip()) < 30 for _, text in pages)
    warnings = []
    if suffix == ".pdf":
        warnings.append("当前索引只包含 PDF 文字层，图片中的图表、公式、方框缺字及重排版面请核验原 PDF；尚未自动 OCR 或理解图像")
    if blank:
        warnings.append(f"{blank} 页没有可提取文字，图片/扫描页尚未 OCR")
    if short:
        warnings.append(f"{short} 页文字少于 30 字，请核验是否为图表、封面或识别不足")
    pages.quality.update(total_pages=len(pages), text_pages=len(pages) - blank, blank_pages=blank, short_pages=short,
                         character_count=sum(len(text) for _, text in pages), warnings=warnings,
                         page_kind="physical" if suffix == ".pdf" else "logical")
    if progress:
        progress({"stage": "chunking", "current_page": len(pages), "total_pages": len(pages), "chunk_count": 0, "embedding_percent": 0})
    return pages


def logical_pages(text, limit=1800):
    pages, buffer = [], ""
    for paragraph in text.replace("\r\n", "\n").split("\n"):
        if len(buffer) + len(paragraph) > limit and buffer:
            pages.append((len(pages) + 1, buffer))
            buffer = ""
        while len(paragraph) > limit:
            if buffer:
                pages.append((len(pages) + 1, buffer))
                buffer = ""
            pages.append((len(pages) + 1, paragraph[:limit]))
            paragraph = paragraph[limit:]
        buffer += paragraph + "\n"
    if buffer.strip():
        pages.append((len(pages) + 1, buffer))
    return pages


def chunk_pages(pages, limit=850, overlap=120, progress=None):
    chunks = []
    section = "正文"
    parent = ""
    # A genuine PDF outline already supplies the hierarchy. A printed running
    # header/footer must not overwrite it with a new chapter or a page suffix.
    outline = getattr(pages, "quality", {}).get("pdf_outline", [])
    outline_labels = {}
    for entry in outline:
        path = entry["path"]
        key = re.sub(r"\s+", "", path.rsplit(" / ", 1)[-1])
        outline_labels.setdefault(key, []).append(path)
    for page, raw in pages:
        mapped = getattr(pages, "sections", {}).get(page)
        if mapped and mapped != "正文":
            section = mapped
            parent = mapped.split(" / ")[0]
        raw = re.sub(r"[ \t]+", " ", raw).strip()
        physical = getattr(pages, "quality", {}).get("page_kind") == "physical"
        if physical:
            headings = [match for match in re.finditer(r"(?m)^[ \t]*(第[一二三四五六七八九十百0-9]+[章节][^\n]{0,65}|\d{1,2}(?:\.\d{1,2}){1,3}[^\n]{2,65})$", raw)
                        if is_pdf_heading(match.group(1).strip())]
        else:
            headings = list(re.finditer(r"(?m)^\s*(?:(#{1,5})\s+([^\n]+)|((?:第[一二三四五六七八九十0-9]+[章节]|\d+[.、]\s*)[^\n]{2,60}))", raw))
        segments = []
        cursor = 0
        for index, heading in enumerate(headings):
            if heading.start() > cursor:
                segments.append((section, raw[cursor:heading.start()]))
            if physical:
                label = heading.group(1).strip()
                if outline:
                    # Restrict same-name matches to this actual outline root;
                    # numbered descendants retain all observed parent levels.
                    key = re.sub(r"\s+", "", label)
                    paths = [path for path in outline_labels.get(key, [])
                             if path.split(" / ", 1)[0] == parent]
                    if len(paths) == 1:
                        section = paths[0]
                    elif re.match(r"^\d+(?:\.\d+)+", label):
                        plain = re.sub(r"\s+\d+\s*$", "", label)
                        known_header = any(path.split(" / ", 1)[0] == parent
                                           for path in outline_labels.get(re.sub(r"\s+", "", plain), []))
                        if not known_header:
                            # Coarse/incomplete bookmarks must not erase an
                            # actual body heading omitted from the PDF outline.
                            section = f"{parent} / {label}" if parent else label
                elif re.match(r"^第[一二三四五六七八九十百0-9]+章", label):
                    parent, section = label, label
                else:
                    section = f"{parent} / {label}" if parent else label
            else:
                label = (heading.group(2) or heading.group(3)).strip()
                depth = len(heading.group(1) or "##")
                if depth <= 2:
                    parent, section = label, label
                else:
                    section = f"{parent} / {label}" if parent else label
            cursor = heading.start()
            end = headings[index + 1].start() if index + 1 < len(headings) else len(raw)
            segments.append((section, raw[cursor:end]))
            cursor = end
        if not headings:
            segments = [(section, raw)]
        for current_section, body in segments:
            start = 0
            while start < len(body):
                end = min(start + limit, len(body))
                if end < len(body):
                    split = max(body.rfind("\n", start + limit // 2, end), body.rfind("。", start + limit // 2, end))
                    if split > start:
                        end = split + 1
                text = body[start:end].strip()
                if text:
                    chunks.append({"id": uid(), "page": page, "section": current_section, "text": text})
                if end >= len(body):
                    break
                start = max(start + 1, end - overlap)
        if progress and (page % 10 == 0 or page == len(pages)):
            progress({"stage": "chunking", "current_page": page, "total_pages": len(pages), "chunk_count": len(chunks), "embedding_percent": 0})
    return chunks


def rebuild_extracted_graph(conn, course_id):
    """Refresh the legacy heading sample from the same complete read model."""
    course = conn.execute("SELECT is_builtin FROM courses WHERE id=?", (course_id,)).fetchone()
    if not course or course["is_builtin"]:
        return
    documents = conn.execute("SELECT id,title,source_type,page_count,created_at,quality FROM documents "
                             "WHERE course_id=? AND status='ready' ORDER BY created_at,id", (course_id,)).fetchall()
    rows = conn.execute("SELECT c.document_id,c.page,c.section,c.text FROM chunks c JOIN documents d ON d.id=c.document_id "
                        "WHERE c.course_id=? AND d.status='ready' ORDER BY d.created_at,d.id,c.page,c.rowid", (course_id,)).fetchall()
    nodes, edges = extracted_graph(derive_textbook_tree(documents, rows))
    conn.execute("DELETE FROM concepts WHERE course_id=?", (course_id,))
    conn.execute("DELETE FROM edges WHERE course_id=?", (course_id,))
    for node in nodes:
        conn.execute("INSERT INTO concepts VALUES (?,?,?,?,?)", (
            node["id"], course_id, node["label"], node["chapter"], node["source"]))
    conn.executemany("INSERT OR IGNORE INTO edges VALUES (?,?,?,?)", [
        (course_id, edge["source"], edge["target"], edge["relation"]) for edge in edges])


def store_chunks(db, doc, pages, chunks, vectors=None, embedding_status="lexical"):
    with db.connection() as conn:
        # Deletion while ingestion runs must never resurrect a removed document.
        if not conn.execute("SELECT 1 FROM documents WHERE id=?", (doc["id"],)).fetchone():
            return False
        conn.execute("DELETE FROM chunks WHERE document_id=?", (doc["id"],))
        conn.executemany("INSERT INTO chunks VALUES (?,?,?,?,?,?,?,?)", [
            (c["id"], doc["id"], doc["course_id"], doc["user_id"], c["page"], c["section"], c["text"],
             dumps(vectors[i]) if vectors else None)
            for i, c in enumerate(chunks)
        ])
        quality = dict(getattr(pages, "quality", {}))
        quality["ingestion_version"] = INGESTION_VERSION
        quality.update(semantic_status=embedding_status, embedded_chunks=len(vectors) if vectors else 0)
        progress = {"stage": "ready", "current_page": len(pages), "total_pages": len(pages), "chunk_count": len(chunks),
                    "embedding_percent": 100 if vectors else 0, "embedding_completed": len(vectors) if vectors else 0, "updated_at": now()}
        conn.execute("UPDATE documents SET status='ready',page_count=?,chunk_count=?,error=NULL,embedding_status=?,progress=?,quality=? WHERE id=?",
                     (len(pages), len(chunks), embedding_status, dumps(progress), dumps(quality), doc["id"]))
        conn.execute("UPDATE ingestion_jobs SET status='ready',error=NULL,updated_at=? WHERE document_id=?", (now(), doc["id"]))
        if getattr(pages, "figures", None) is not None:
            store_figure_index(conn, doc, pages.figures, len(pages))
        rebuild_extracted_graph(conn, doc["course_id"])
    return True


def seed_catalog(db, settings):
    path = settings.seed_file
    if not path.exists():
        return
    catalog = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(catalog, dict):
        catalog = catalog.get("courses", [])
    for course in catalog:
        course_id = course["id"]
        with db.connection() as conn:
            conn.execute("INSERT INTO courses VALUES (?,?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET "
                         "title=excluded.title,description=excluded.description,code=excluded.code,color=excluded.color,"
                         "icon=excluded.icon,assistant_name=excluded.assistant_name", (
                             course_id, None, course["title"], course["description"], course["code"],
                             course["color"], course["icon"], course["assistant_name"], 1, now()))
            conn.execute("DELETE FROM concepts WHERE course_id=?", (course_id,))
            conn.executemany("INSERT INTO concepts VALUES (?,?,?,?,?)", [
                (c["id"], course_id, c["label"], c["chapter"], "课程手册知识结构") for c in course.get("concepts", [])])
            conn.execute("DELETE FROM edges WHERE course_id=?", (course_id,))
            conn.executemany("INSERT INTO edges VALUES (?,?,?,?)", [
                (course_id, e["source"], e["target"], e["relation"]) for e in course.get("edges", [])])
            conn.execute("DELETE FROM question_bank WHERE course_id=?", (course_id,))
            conn.executemany("INSERT INTO question_bank VALUES (?,?,?)", [
                (q["id"], course_id, dumps(q)) for q in course.get("questions", [])])
        for source in course.get("documents", []):
            file_path = (settings.root / source["path"]).resolve()
            if not file_path.is_relative_to(settings.root.resolve()) or not file_path.is_file():
                raise ValueError(f"种子教材路径无效：{source['path']}")
            digest = hashlib.sha256(INGESTION_VERSION.encode() + b"\0" + file_path.read_bytes()).hexdigest()
            doc_id = "seed-" + hashlib.sha256(f"{course_id}/{source['path']}".encode()).hexdigest()[:24]
            existing = db.one("SELECT * FROM documents WHERE id=?", (doc_id,))
            if existing and existing["seed_hash"] == digest and existing["status"] == "ready":
                continue
            db.execute("INSERT INTO documents (id,course_id,user_id,title,filename,path,status,source_type,license,source_url,created_at,seed_hash) "
                       "VALUES (?,?,?,?,?,?,'processing',?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET seed_hash=excluded.seed_hash,path=excluded.path,"
                       "title=excluded.title,license=excluded.license,source_url=excluded.source_url,source_type=excluded.source_type", (
                           doc_id, course_id, None, source["title"], file_path.name, str(file_path),
                           source.get("source_type", "original_handbook"), source["license"], source.get("source_url"), now(), digest))
            doc = db.one("SELECT * FROM documents WHERE id=?", (doc_id,))
            try:
                pages = extract_pages(file_path)
                store_chunks(db, doc, pages, chunk_pages(pages))
            except Exception:
                db.execute("UPDATE documents SET status='failed',error=? WHERE id=?", ("种子教材无法读取，请检查源文件", doc_id))


class IngestionWorker:
    def __init__(self, db, provider, settings):
        self.db, self.provider, self.settings = db, provider, settings
        self.task = None
        self.stopping = False

    async def start(self):
        # A killed process leaves processing jobs; requeue them and replace chunks atomically.
        self.db.execute("UPDATE ingestion_jobs SET status='queued',updated_at=? WHERE status='processing'", (now(),))
        self.db.execute("UPDATE documents SET status='queued' WHERE id IN (SELECT document_id FROM ingestion_jobs WHERE status='queued')")
        queue_missing_figure_indexes(self.db)
        self.stopping = False
        self.task = asyncio.create_task(self.run())

    async def stop(self):
        self.stopping = True
        if self.task:
            self.task.cancel()
            try:
                await self.task
            except asyncio.CancelledError:
                pass

    async def run(self):
        while not self.stopping:
            job = self.db.one("SELECT j.id AS job_id,d.* FROM ingestion_jobs j JOIN documents d ON d.id=j.document_id "
                              "WHERE j.status='queued' ORDER BY j.created_at LIMIT 1")
            if not job:
                figure_job = self.db.one("SELECT d.* FROM figure_index_jobs f JOIN documents d ON d.id=f.document_id "
                                         "WHERE f.status='queued' AND d.status='ready' ORDER BY f.updated_at,d.id LIMIT 1")
                if figure_job:
                    await self.process_figures(figure_job)
                    continue
                await asyncio.sleep(0.4)
                continue
            try:
                await self.process(job)
            except asyncio.CancelledError:
                raise
            except Exception:
                self.fail(job, "教材摄取遇到内部错误，请重新上传")

    def fail(self, doc, error):
        with self.db.connection() as conn:
            row = conn.execute("SELECT progress FROM documents WHERE id=?", (doc["id"],)).fetchone()
            progress = loads(row["progress"], {}) if row else {}
            progress.update(stage="failed", updated_at=now())
            conn.execute("UPDATE documents SET status='failed',error=?,progress=? WHERE id=?", (error, dumps(progress), doc["id"]))
            conn.execute("UPDATE ingestion_jobs SET status='failed',error=?,updated_at=? WHERE document_id=?", (error, now(), doc["id"]))

    def update_progress(self, doc_id, value):
        value = {**value, "updated_at": now()}
        self.db.execute("UPDATE documents SET progress=? WHERE id=?", (dumps(value), doc_id))

    def stage_chunks(self, doc, pages, chunks):
        with self.db.connection() as conn:
            if not conn.execute("SELECT 1 FROM documents WHERE id=?", (doc["id"],)).fetchone():
                return False
            conn.execute("DELETE FROM chunks WHERE document_id=?", (doc["id"],))
            conn.executemany("INSERT INTO chunks VALUES (?,?,?,?,?,?,?,NULL)", [
                (c["id"], doc["id"], doc["course_id"], doc["user_id"], c["page"], c["section"], c["text"]) for c in chunks])
            progress = {"stage": "embedding", "current_page": len(pages), "total_pages": len(pages), "chunk_count": len(chunks),
                        "embedding_completed": 0, "embedding_percent": 0, "updated_at": now()}
            conn.execute("UPDATE documents SET page_count=?,chunk_count=?,quality=?,progress=? WHERE id=?", (
                len(pages), len(chunks), dumps({**getattr(pages, "quality", {}), "ingestion_version": INGESTION_VERSION}), dumps(progress), doc["id"]))
            if getattr(pages, "figures", None) is not None:
                store_figure_index(conn, doc, pages.figures, len(pages))
        return True

    async def process_figures(self, doc):
        """Backfill only caption metadata; keep ready chunks and paid vectors intact."""
        with self.db.connection() as conn:
            claimed = conn.execute("UPDATE figure_index_jobs SET status='processing',attempts=attempts+1,error=NULL,updated_at=? "
                                   "WHERE document_id=? AND status='queued' AND EXISTS "
                                   "(SELECT 1 FROM documents WHERE id=? AND status='ready')",
                                   (now(), doc["id"], doc["id"])).rowcount
        if not claimed:
            return
        try:
            rows = self.db.all("SELECT page,section FROM chunks WHERE document_id=? ORDER BY page,rowid", (doc["id"],))
            candidates = {}
            for row in rows:
                candidates.setdefault(row["page"], set()).add(row["section"])
            sections = {page: next(iter(values)) for page, values in candidates.items() if len(values) == 1}

            def report(current, total):
                self.db.execute("UPDATE figure_index_jobs SET current_page=?,total_pages=?,updated_at=? "
                                "WHERE document_id=? AND status='processing'", (current, total, now(), doc["id"]))

            figures = await asyncio.to_thread(scan_pdf_figures, doc["path"], doc["page_count"], sections, report)
            with self.db.connection() as conn:
                # A removed or newly retrying document must not be resurrected.
                current = conn.execute("SELECT * FROM documents WHERE id=? AND status='ready'", (doc["id"],)).fetchone()
                if current:
                    store_figure_index(conn, current, figures, doc["page_count"])
                else:
                    conn.execute("UPDATE figure_index_jobs SET status='queued',updated_at=? WHERE document_id=?", (now(), doc["id"]))
        except asyncio.CancelledError:
            # The persisted processing checkpoint is requeued at the next start.
            raise
        except Exception:
            self.db.execute("UPDATE figure_index_jobs SET status='failed',error=?,updated_at=? WHERE document_id=?",
                            ("PDF 图题定位失败；原文字索引可继续使用，请重试或核验原文件", now(), doc["id"]))

    def finish(self, doc, quality, semantic_status, total_chunks, embedded):
        with self.db.connection() as conn:
            if not conn.execute("SELECT 1 FROM documents WHERE id=?", (doc["id"],)).fetchone():
                return
            quality.update(semantic_status=semantic_status, embedded_chunks=embedded)
            progress = {"stage": "ready", "current_page": doc["page_count"], "total_pages": doc["page_count"], "chunk_count": total_chunks,
                        "embedding_completed": embedded, "embedding_percent": round(100 * embedded / max(total_chunks, 1), 1), "updated_at": now()}
            conn.execute("UPDATE documents SET status='ready',error=NULL,embedding_status=?,quality=?,progress=? WHERE id=?", (
                semantic_status, dumps(quality), dumps(progress), doc["id"]))
            conn.execute("UPDATE ingestion_jobs SET status='ready',error=NULL,updated_at=? WHERE document_id=?", (now(), doc["id"]))
            ready = conn.execute("SELECT id,status,source_type,page_count FROM documents WHERE id=?", (doc["id"],)).fetchone()
            queue_figure_index(conn, ready)
            rebuild_extracted_graph(conn, doc["course_id"])

    async def process(self, doc):
        with self.db.connection() as conn:
            conn.execute("UPDATE ingestion_jobs SET status='processing',attempts=attempts+1,updated_at=? WHERE document_id=?", (now(), doc["id"]))
            conn.execute("UPDATE documents SET status='processing' WHERE id=?", (doc["id"],))
        try:
            previous_progress = loads(doc.get("progress"), {})
            quality = loads(doc.get("quality"), {})
            # Persisted vectors can resume after restart or a transient embedding failure.
            chunks = self.db.all("SELECT id,page,section,text,embedding FROM chunks WHERE document_id=? ORDER BY page,rowid", (doc["id"],))
            resumable = bool(chunks) and (previous_progress.get("stage") in {"embedding", "persisting"} or quality.get("semantic_status") == "lexical_fallback")
            if not resumable:
                report = lambda value: self.update_progress(doc["id"], value)
                pages = await asyncio.to_thread(extract_pages, Path(doc["path"]), report)
                chunks = await asyncio.to_thread(chunk_pages, pages, 850, 120, report)
                if not self.stage_chunks(doc, pages, chunks):
                    return
                doc = self.db.one("SELECT * FROM documents WHERE id=?", (doc["id"],))
                if not doc:
                    return
                quality = {**pages.quality, "ingestion_version": INGESTION_VERSION}
            embedding_status = "lexical"
            embedded = sum(bool(c.get("embedding")) for c in chunks)
            quality.pop("semantic_error", None)
            if self.settings.embedding_enabled and (self.settings.embedding_api_key or self.settings.api_key):
                try:
                    pending = [c for c in chunks if not c.get("embedding")]
                    for start in range(0, len(pending), 32):
                        if not self.db.one("SELECT 1 FROM documents WHERE id=?", (doc["id"],)):
                            return
                        current = pending[start:start + 32]
                        self.update_progress(doc["id"], {"stage": "embedding", "current_page": doc["page_count"], "total_pages": doc["page_count"], "chunk_count": len(chunks),
                                                          "embedding_completed": embedded, "embedding_percent": round(100 * embedded / max(len(chunks), 1), 1)})
                        batch = await self.provider.embed([c["text"] for c in current])
                        if len(batch) != len(current):
                            raise ValueError("向量数量不匹配")
                        with self.db.connection() as conn:
                            conn.executemany("UPDATE chunks SET embedding=? WHERE id=? AND document_id=?", [(dumps(vector), chunk["id"], doc["id"]) for chunk, vector in zip(current, batch)])
                        embedded += len(batch)
                        self.update_progress(doc["id"], {"stage": "embedding", "current_page": doc["page_count"], "total_pages": doc["page_count"], "chunk_count": len(chunks),
                                                          "embedding_completed": embedded, "embedding_percent": round(100 * embedded / max(len(chunks), 1), 1)})
                    embedding_status = "hybrid"
                    quality["warnings"] = [warning for warning in quality.get("warnings", []) if not warning.startswith("语义向量尚未完整生成")]
                except Exception as exc:
                    from .provider import ProviderError
                    embedding_status = "lexical_fallback"
                    quality["semantic_error"] = str(exc) if isinstance(exc, ProviderError) else "语义向量生成失败，请重试；当前仍可使用本地词法检索"
                    warning = "语义向量尚未完整生成，当前教材使用词法检索或部分向量；可点击重试补全"
                    quality.setdefault("warnings", [])
                    if warning not in quality["warnings"]:
                        quality["warnings"].append(warning)
            else:
                quality["semantic_note"] = "语义向量未启用或未提供服务器凭证，当前仅本地词法检索"
            self.update_progress(doc["id"], {"stage": "persisting", "current_page": doc["page_count"], "total_pages": doc["page_count"], "chunk_count": len(chunks),
                                               "embedding_completed": embedded, "embedding_percent": round(100 * embedded / max(len(chunks), 1), 1)})
            self.finish(doc, quality, embedding_status, len(chunks), embedded)
        except ValueError as exc:
            self.fail(doc, str(exc))
        except Exception:
            self.fail(doc, "文件内容损坏或格式不匹配，请重新导出后上传")

